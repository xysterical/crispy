from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from app.core.config import get_settings
from app.services.video_frames import sample_video_frames


MAX_FILES = 10
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024
MAX_TOTAL_SIZE_BYTES = 200 * 1024 * 1024

SKU_EXTENSIONS = {".csv", ".xlsx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v"}


def _safe_name(name: str) -> str:
    return Path(name).name.replace(" ", "_")


def _allocate_unique_token(token: str, seen: set[str]) -> str:
    candidate = token
    idx = 2
    while candidate in seen:
        candidate = f"{token}_{idx}"
        idx += 1
    seen.add(candidate)
    return candidate


def _allocate_unique_filename(name: str, seen: set[str]) -> str:
    path = Path(name)
    stem = path.stem or "upload"
    suffix = path.suffix
    token = _allocate_unique_token(stem, {Path(item).stem for item in seen})
    candidate = f"{token}{suffix}"
    while candidate in seen:
        token = _allocate_unique_token(stem, {Path(item).stem for item in seen})
        candidate = f"{token}{suffix}"
    seen.add(candidate)
    return candidate


def _parse_csv_bytes(data: bytes) -> list[dict]:
    text = data.decode("utf-8", errors="ignore")
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict] = []
    for idx, row in enumerate(reader):
        if idx >= 200:
            break
        rows.append({k: (v or "").strip() for k, v in row.items()})
    return rows


def _parse_xlsx_bytes(data: bytes) -> list[dict]:
    try:
        import openpyxl  # type: ignore
    except Exception:
        raise RuntimeError("openpyxl_not_installed")
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(min_row=1, max_row=201, values_only=True))
    if not rows:
        return []
    headers = [str(item) if item is not None else "" for item in rows[0]]
    parsed: list[dict] = []
    for row in rows[1:]:
        parsed.append({headers[idx]: row[idx] for idx in range(min(len(headers), len(row)))})
    return parsed


def _extract_video_frames(run_id: str, variant: str, video_path: Path, count: int = 3) -> list[str]:
    settings = get_settings()
    frame_dir = settings.assets_dir / run_id / "input_frames"
    return sample_video_frames(
        video_path=video_path,
        output_dir=frame_dir,
        prefix=variant,
        count=count,
    )


def process_uploaded_payloads(run_id: str, uploads: list[dict[str, Any]]) -> tuple[dict, list[dict]]:
    if len(uploads) > MAX_FILES:
        raise ValueError(f"too many files; max={MAX_FILES}")

    settings = get_settings()
    input_dir = settings.assets_dir / run_id / "input"
    input_dir.mkdir(parents=True, exist_ok=True)

    total_size = 0
    sku_summary: list[dict] = []
    sample_images: list[dict] = []
    sample_videos: list[dict] = []
    file_entries: list[dict] = []
    artifacts: list[dict] = []
    stored_filenames: set[str] = set()
    video_prefixes: set[str] = set()

    for item in uploads:
        filename = _safe_name(item["filename"])
        stored_filename = _allocate_unique_filename(filename, stored_filenames)
        content = item["content"]
        content_type = item.get("content_type")
        size = len(content)
        total_size += size
        if size > MAX_FILE_SIZE_BYTES:
            raise ValueError(f"file too large: {filename} max={MAX_FILE_SIZE_BYTES}")
        if total_size > MAX_TOTAL_SIZE_BYTES:
            raise ValueError(f"total upload too large; max={MAX_TOTAL_SIZE_BYTES}")

        path = input_dir / stored_filename
        path.write_bytes(content)
        ext = path.suffix.lower()
        entry = {
            "filename": filename,
            "content_type": content_type,
            "size_bytes": size,
            "uri": str(path),
            "extension": ext,
        }
        file_entries.append(entry)
        artifacts.append({"type": "input_file", "uri": str(path), "payload": entry})

        if ext in SKU_EXTENSIONS:
            try:
                parsed_rows = _parse_csv_bytes(content) if ext == ".csv" else _parse_xlsx_bytes(content)
                sku_summary.extend(parsed_rows[:100])
            except Exception as exc:
                sku_summary.append({"filename": filename, "parse_error": str(exc)})
        elif ext in IMAGE_EXTENSIONS:
            sample_images.append(entry)
        elif ext in VIDEO_EXTENSIONS:
            frame_prefix = _allocate_unique_token(Path(filename).stem or "video", video_prefixes)
            frames = _extract_video_frames(run_id, frame_prefix, path, count=3)
            sample_videos.append({**entry, "frame_placeholders": frames, "duration_seconds": 15.0})
            for frame_path in frames:
                artifacts.append(
                    {
                        "type": "input_video_frame",
                        "uri": frame_path,
                        "payload": {"source_video": str(path), "frame_uri": frame_path},
                    }
                )

    summary = {
        "files": file_entries,
        "sku_summary": sku_summary,
        "sample_images": sample_images,
        "sample_videos": sample_videos,
    }
    return summary, artifacts
